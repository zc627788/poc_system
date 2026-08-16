from __future__ import annotations

import hashlib
import json
import re
import secrets
import string
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models.billing_order import BillingOrder
from app.models.consumption_import import ConsumptionImportBatch, ConsumptionImportRow
from app.models.consumption_record import ConsumptionRecord
from app.models.prepaid_ledger_entry import PrepaidLedgerEntry
from app.repositories.prepaid_billing import (
    get_consumption_import_batch,
    get_consumption_record_by_unique_key,
    get_prepaid_ledger_entry_by_source,
    list_admin_balances,
    list_admin_consumptions,
    list_consumption_import_rows,
    list_import_rows_by_status,
    list_orders_by_order_nos,
    list_paid_prepaid_orders,
    list_public_balances,
    list_public_consumptions,
    list_service_reference_products,
    list_site_users_by_phones,
)
from app.schemas.prepaid_billing import (
    AdminBalanceListResponse,
    AdminBalanceSummaryResponse,
    BalanceSummaryResponse,
    ConsumptionImportBatchDetailResponse,
    ConsumptionImportBatchStatus,
    ConsumptionImportBatchSummaryResponse,
    ConsumptionImportConfirmResponse,
    ConsumptionImportPreviewResponse,
    ConsumptionImportRowResponse,
    ConsumptionImportRowStatus,
    ConsumptionRecordListResponse,
    ConsumptionRecordResponse,
    ConsumptionSourceRecord,
    ConsumptionSummaryResponse,
    PrepaidBackfillResultResponse,
    PrepaidLedgerSourceType,
)



UTC = timezone.utc
SHANGHAI_TZ = timezone(timedelta(hours=8), name="Asia/Shanghai")
MAX_IMPORT_FILE_BYTES = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 10_000
PREVIEW_PAGE_SIZE = 100
IMPORT_BATCH_ALPHABET = string.ascii_uppercase + string.digits
PHONE_PATTERN = re.compile(r"^1[3-9]\d{9}$")
CODE_PATTERN = re.compile(r"^[A-Z0-9_]{2,64}$")
EXCEL_HEADERS = [
    "供应商流水号",
    "消费时间",
    "用户手机号",
    "服务编码",
    "供应商名称",
    "消费项目",
    "消费数量",
    "用量单位",
    "实际消费金额",
    "关联订单号",
    "备注",
]


class PrepaidBillingValidationError(ValueError):
    pass


class PrepaidBillingConflictError(ValueError):
    pass


class PrepaidBillingNotFoundError(ValueError):
    pass


def ensure_prepaid_order_credit_entry(
    db: Session,
    order: BillingOrder,
    occurred_at: datetime | None = None,
) -> bool:
    if order.billing_mode != "PREPAID" or order.status != "PAID":
        return False
    inserted = False
    existing = get_prepaid_ledger_entry_by_source(db, PrepaidLedgerSourceType.PAID_ORDER.value, order.id, for_update=True)
    if existing is None:
        ledger = PrepaidLedgerEntry(
            site_user_id=order.site_user_id,
            service_code=order.service_code,
            service_name=order.service_name,
            entry_type="CREDIT",
            amount=Decimal(order.order_amount).quantize(Decimal("0.01")),
            currency=order.currency,
            source_type=PrepaidLedgerSourceType.PAID_ORDER.value,
            source_id=order.id,
            occurred_at=_naive_utc(occurred_at) or order.paid_at or datetime.now(UTC).replace(tzinfo=None),
        )
        try:
            with db.begin_nested():
                db.add(ledger)
                db.flush()
            inserted = True
        except IntegrityError:
            inserted = False
    return inserted


def backfill_paid_prepaid_orders(db: Session, dry_run: bool = False) -> PrepaidBackfillResultResponse:
    inserted = 0
    orders = list_paid_prepaid_orders(db)
    for order in orders:
        if ensure_prepaid_order_credit_entry(db, order, _utc_value(order.paid_at)):
            inserted += 1
    if dry_run:
        db.rollback()
    else:
        db.commit()
    return PrepaidBackfillResultResponse(scanned_orders=len(orders), inserted_entries=inserted)


def build_consumption_import_template() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "消费导入模板"
    worksheet.append(EXCEL_HEADERS)
    worksheet.append(
        [
            "SUP-20260714-0001",
            "2026-07-14 10:00:00",
            "13900009999",
            "AI_COMPUTE",
            "创新算力服务商",
            "GPU 推理调用",
            "12.5",
            "次",
            "88.80",
            "",
            "示例数据，请删除后再导入",
        ]
    )
    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 14), 28)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()


def build_consumption_import_errors_workbook(rows: list[ConsumptionImportRow]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "导入错误明细"
    worksheet.append([*EXCEL_HEADERS, "错误原因"])
    for row in rows:
        raw_values = _decode_raw_values(row.raw_values_json)
        worksheet.append([raw_values.get(header) or "" for header in EXCEL_HEADERS] + [row.error_message or ""])
    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 14), 36)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()


def preview_consumption_import(
    db: Session,
    admin_user_id: int,
    file_name: str,
    file_bytes: bytes,
) -> ConsumptionImportPreviewResponse:
    normalized_file_name = (file_name or "").strip()
    if not normalized_file_name.lower().endswith(".xlsx"):
        raise PrepaidBillingValidationError("仅支持上传 .xlsx 文件")
    if len(file_bytes) > MAX_IMPORT_FILE_BYTES:
        raise PrepaidBillingValidationError("上传文件不能超过 10 MB")

    records = _parse_excel_source_records(file_bytes)
    uploaded_at = datetime.now(UTC).replace(tzinfo=None)
    batch = ConsumptionImportBatch(
        batch_no=_generate_import_batch_no(uploaded_at.replace(tzinfo=UTC)),
        record_source="EXCEL",
        file_name=normalized_file_name,
        file_sha256=hashlib.sha256(file_bytes).hexdigest(),
        status=ConsumptionImportBatchStatus.PARSING.value,
        total_rows=len(records),
        uploaded_by=admin_user_id,
        uploaded_at=uploaded_at,
        expires_at=uploaded_at + timedelta(hours=24),
        created_at=uploaded_at,
        updated_at=uploaded_at,
    )
    db.add(batch)
    db.flush()

    rows = _build_preview_rows(db, batch.id, records, uploaded_at)
    db.add_all(rows)
    batch.status = ConsumptionImportBatchStatus.PREVIEWED.value
    batch.valid_rows = sum(1 for item in rows if item.row_status == ConsumptionImportRowStatus.VALID.value)
    batch.invalid_rows = sum(1 for item in rows if item.row_status == ConsumptionImportRowStatus.INVALID.value)
    batch.updated_at = uploaded_at
    db.commit()
    return get_consumption_import_batch_detail(db, batch.id, None, 1, PREVIEW_PAGE_SIZE)


def confirm_consumption_import(
    db: Session,
    batch_id: int,
    admin_user_id: int,
) -> ConsumptionImportConfirmResponse:
    batch = get_consumption_import_batch(db, batch_id, for_update=True)
    if batch is None:
        raise PrepaidBillingNotFoundError("导入批次不存在")
    _expire_batch_if_needed(db, batch, raise_conflict=True)
    if batch.status == ConsumptionImportBatchStatus.COMPLETED.value:
        return _batch_summary_response(batch)
    if batch.status != ConsumptionImportBatchStatus.PREVIEWED.value:
        raise PrepaidBillingConflictError("当前批次状态不允许确认导入")

    valid_rows = list_import_rows_by_status(db, batch.id, ConsumptionImportRowStatus.VALID.value, for_update=True)
    if not valid_rows:
        raise PrepaidBillingValidationError("当前批次没有可导入的有效记录")

    now = datetime.now(UTC).replace(tzinfo=None)
    batch.status = ConsumptionImportBatchStatus.IMPORTING.value
    batch.updated_at = now
    db.flush()

    imported_rows = 0
    try:
        for row in valid_rows:
            if _import_consumption_row(db, row):
                imported_rows += 1
        batch.status = ConsumptionImportBatchStatus.COMPLETED.value
        batch.confirmed_by = admin_user_id
        batch.confirmed_at = now
        batch.imported_rows = imported_rows
        batch.valid_rows = db.query(ConsumptionImportRow).filter(
            ConsumptionImportRow.batch_id == batch.id,
            ConsumptionImportRow.row_status == ConsumptionImportRowStatus.VALID.value,
        ).count()
        batch.invalid_rows = db.query(ConsumptionImportRow).filter(
            ConsumptionImportRow.batch_id == batch.id,
            ConsumptionImportRow.row_status == ConsumptionImportRowStatus.INVALID.value,
        ).count()
        batch.updated_at = now
        db.commit()
    except Exception as exc:
        db.rollback()
        locked = get_consumption_import_batch(db, batch_id, for_update=True)
        if locked is not None:
            locked.status = ConsumptionImportBatchStatus.FAILED.value
            locked.error_message = str(exc)[:500]
            locked.updated_at = datetime.now(UTC).replace(tzinfo=None)
            db.commit()
        raise

    refreshed = get_consumption_import_batch(db, batch_id)
    if refreshed is None:
        raise PrepaidBillingNotFoundError("导入批次不存在")
    return _batch_summary_response(refreshed)


def get_consumption_import_batch_detail(
    db: Session,
    batch_id: int,
    row_status: ConsumptionImportRowStatus | None,
    page: int,
    page_size: int,
) -> ConsumptionImportBatchDetailResponse:
    batch = get_consumption_import_batch(db, batch_id, for_update=False)
    if batch is None:
        raise PrepaidBillingNotFoundError("导入批次不存在")
    _expire_batch_if_needed(db, batch, raise_conflict=False)
    rows, total = list_consumption_import_rows(db, batch.id, page, page_size, row_status.value if row_status else None)
    return ConsumptionImportBatchDetailResponse(
        batch=_batch_summary_response(batch),
        items=[_row_response(item) for item in rows],
        total=total,
        page=page,
        page_size=page_size,
    )


def list_consumption_import_errors(db: Session, batch_id: int) -> list[ConsumptionImportRow]:
    batch = get_consumption_import_batch(db, batch_id)
    if batch is None:
        raise PrepaidBillingNotFoundError("导入批次不存在")
    _expire_batch_if_needed(db, batch, raise_conflict=False)
    return list_import_rows_by_status(db, batch.id, ConsumptionImportRowStatus.INVALID.value)


def get_public_balance_summaries(db: Session, site_user_id: int) -> list[BalanceSummaryResponse]:
    return [
        BalanceSummaryResponse(
            service_code=service_code,
            service_name=service_name,
            total_recharged=Decimal(total_recharged).quantize(Decimal("0.01")),
            total_consumed=Decimal(total_consumed).quantize(Decimal("0.01")),
            remaining_balance=(Decimal(total_recharged) - Decimal(total_consumed)).quantize(Decimal("0.01")),
            insufficient=(Decimal(total_recharged) - Decimal(total_consumed)) < 0,
            last_updated_at=_utc_value(last_updated_at),
        )
        for service_code, service_name, total_recharged, total_consumed, last_updated_at in list_public_balances(db, site_user_id)
    ]


def get_public_consumption_records(
    db: Session,
    site_user_id: int,
    page: int,
    page_size: int,
    service_code: str | None = None,
    start_at: datetime | None = None,
    end_at: datetime | None = None,
    order_no: str | None = None,
) -> ConsumptionRecordListResponse:
    items, total, total_amount = list_public_consumptions(db, site_user_id, page, page_size, service_code, start_at, end_at, order_no)
    return ConsumptionRecordListResponse(
        items=[
            _consumption_record_response(record, None, None, linked_order_no)
            for record, linked_order_no in items
        ],
        total=total,
        page=page,
        page_size=page_size,
        summary=ConsumptionSummaryResponse(total_amount=Decimal(total_amount).quantize(Decimal("0.01"))),
    )


def get_admin_balance_summaries(
    db: Session,
    page: int,
    page_size: int,
    user_keyword: str | None = None,
    service_code: str | None = None,
    insufficient: bool | None = None,
) -> AdminBalanceListResponse:
    items, total = list_admin_balances(db, page, page_size, user_keyword, service_code, insufficient)
    return AdminBalanceListResponse(
        items=[
            AdminBalanceSummaryResponse(
                site_user_id=item.site_user_id,
                user_name=item.name,
                user_phone_masked=item.mobile_fuzzy or _mask_phone(item.phone),
                service_code=item.service_code,
                service_name=item.service_name,
                total_recharged=Decimal(item.total_recharged).quantize(Decimal("0.01")),
                total_consumed=Decimal(item.total_consumed).quantize(Decimal("0.01")),
                remaining_balance=Decimal(item.remaining_balance).quantize(Decimal("0.01")),
                insufficient=Decimal(item.remaining_balance) < 0,
                last_updated_at=_utc_value(item.last_updated_at),
            )
            for item in items
        ],
        total=total,
        page=page,
        page_size=page_size,
    )


def get_admin_consumption_records(
    db: Session,
    page: int,
    page_size: int,
    user_keyword: str | None = None,
    service_code: str | None = None,
    supplier_name: str | None = None,
    start_at: datetime | None = None,
    end_at: datetime | None = None,
    order_no: str | None = None,
) -> ConsumptionRecordListResponse:
    items, total, total_amount = list_admin_consumptions(
        db, page, page_size, user_keyword, service_code, supplier_name, start_at, end_at, order_no
    )
    return ConsumptionRecordListResponse(
        items=[
            _consumption_record_response(record, user_name, mobile_fuzzy or _mask_phone(phone), linked_order_no)
            for record, user_name, phone, mobile_fuzzy, linked_order_no in items
        ],
        total=total,
        page=page,
        page_size=page_size,
        summary=ConsumptionSummaryResponse(total_amount=Decimal(total_amount).quantize(Decimal("0.01"))),
    )


def _build_preview_rows(
    db: Session,
    batch_id: int,
    records: list[tuple[int, ConsumptionSourceRecord, dict[str, str | None]]],
    now: datetime,
) -> list[ConsumptionImportRow]:
    phones = sorted({(record.user_phone or "").strip() for _, record, _ in records if (record.user_phone or "").strip()})
    service_codes = sorted({(record.service_code or "").strip() for _, record, _ in records if (record.service_code or "").strip()})
    order_nos = sorted({(record.order_no or "").strip() for _, record, _ in records if (record.order_no or "").strip()})
    users_by_phone = {item.phone: item for item in list_site_users_by_phones(db, phones)}
    services_by_code = {item.service_code: item for item in list_service_reference_products(db, service_codes)}
    orders_by_no = {item.order_no: item for item in list_orders_by_order_nos(db, order_nos)}

    seen_records: set[tuple[str, str, str]] = set()
    preview_rows: list[ConsumptionImportRow] = []
    for row_no, record, raw_values in records:
        errors: list[str] = []
        supplier_record_no = _normalize_optional_text(record.supplier_record_no, 128, "供应商流水号", errors)
        user_phone = _normalize_phone(record.user_phone, errors)
        service_code = _normalize_code(record.service_code, "服务编码", errors)
        supplier_name = _normalize_required_text(record.supplier_name, "供应商名称", 128, errors)
        item_name = _normalize_required_text(record.item_name, "消费项目", 200, errors)
        quantity = _normalize_quantity(record.quantity_text, errors)
        unit = _normalize_optional_text(record.unit, 32, "单位", errors)
        amount = _normalize_amount(record.amount_text, errors)
        order_no = _normalize_optional_text(record.order_no, 40, "关联订单号", errors)
        remark = _normalize_optional_text(record.remark, 500, "备注", errors)
        consumed_at = _normalize_consumed_at(record.consumed_at_text, errors)

        if not supplier_record_no:
            errors.append("供应商流水号不能为空")
        if supplier_name and supplier_record_no:
            duplicate_key = (record.record_source.strip().upper() or "EXCEL", supplier_name, supplier_record_no)
            if duplicate_key in seen_records:
                errors.append("同一文件内存在重复的供应商流水")
            else:
                seen_records.add(duplicate_key)
            if get_consumption_record_by_unique_key(db, duplicate_key[0], duplicate_key[1], duplicate_key[2]) is not None:
                errors.append("该供应商流水已导入，不能重复扣减")

        site_user = users_by_phone.get(user_phone) if user_phone else None
        if user_phone and site_user is None:
            errors.append("用户手机号未匹配到站点用户")

        service_product = services_by_code.get(service_code) if service_code else None
        if service_code and service_product is None:
            errors.append("服务编码不存在")

        order = orders_by_no.get(order_no) if order_no else None
        if order_no:
            if order is None:
                errors.append("关联订单号不存在")
            elif order.status != "PAID":
                errors.append("关联订单必须是已支付订单")
            else:
                if site_user and order.site_user_id != site_user.id:
                    errors.append("关联订单不属于该用户")
                if service_code and order.service_code != service_code:
                    errors.append("关联订单服务与本行服务编码不一致")

        preview_rows.append(
            ConsumptionImportRow(
                batch_id=batch_id,
                row_no=row_no,
                record_source=(record.record_source or "EXCEL").strip().upper() or "EXCEL",
                raw_values_json=json.dumps(raw_values, ensure_ascii=False),
                supplier_record_no=supplier_record_no,
                consumed_at=consumed_at,
                user_phone=user_phone,
                service_code=service_code,
                supplier_name=supplier_name,
                item_name=item_name,
                quantity=quantity,
                unit=unit,
                amount=amount,
                order_no=order_no,
                remark=remark,
                site_user_id=site_user.id if site_user else None,
                service_name=service_product.service_name if service_product else None,
                order_id=order.id if order else None,
                row_status=ConsumptionImportRowStatus.INVALID.value if errors else ConsumptionImportRowStatus.VALID.value,
                error_message="；".join(errors)[:500] if errors else None,
                created_at=now,
                updated_at=now,
            )
        )
    return preview_rows


def _import_consumption_row(db: Session, row: ConsumptionImportRow) -> bool:
    if row.row_status != ConsumptionImportRowStatus.VALID.value:
        return False
    duplicate = get_consumption_record_by_unique_key(
        db,
        row.record_source,
        row.supplier_name or "",
        row.supplier_record_no or "",
        for_update=True,
    )
    if duplicate is not None:
        row.row_status = ConsumptionImportRowStatus.INVALID.value
        row.error_message = "该供应商流水已导入，不能重复扣减"
        row.updated_at = datetime.now(UTC).replace(tzinfo=None)
        db.flush()
        return False

    try:
        with db.begin_nested():
            record = ConsumptionRecord(
                batch_id=row.batch_id,
                import_row_id=row.id,
                record_source=row.record_source,
                supplier_name=row.supplier_name or "",
                supplier_record_no=row.supplier_record_no or "",
                site_user_id=row.site_user_id or 0,
                service_code=row.service_code or "",
                service_name=row.service_name or "",
                order_id=row.order_id,
                consumed_at=row.consumed_at or datetime.now(UTC).replace(tzinfo=None),
                item_name=row.item_name or "",
                quantity=row.quantity,
                unit=row.unit,
                amount=Decimal(row.amount or Decimal("0.00")).quantize(Decimal("0.01")),
                currency="CNY",
                remark=row.remark,
                created_at=datetime.now(UTC).replace(tzinfo=None),
            )
            db.add(record)
            db.flush()

            existing_ledger = get_prepaid_ledger_entry_by_source(
                db,
                PrepaidLedgerSourceType.CONSUMPTION.value,
                record.id,
                for_update=True,
            )
            if existing_ledger is None:
                db.add(
                    PrepaidLedgerEntry(
                        site_user_id=record.site_user_id,
                        service_code=record.service_code,
                        service_name=record.service_name,
                        entry_type="DEBIT",
                        amount=record.amount,
                        currency="CNY",
                        source_type=PrepaidLedgerSourceType.CONSUMPTION.value,
                        source_id=record.id,
                        occurred_at=record.consumed_at,
                    )
                )
                db.flush()
    except IntegrityError:
        row.row_status = ConsumptionImportRowStatus.INVALID.value
        row.error_message = "该供应商流水已导入，不能重复扣减"
        row.updated_at = datetime.now(UTC).replace(tzinfo=None)
        db.flush()
        return False

    row.row_status = ConsumptionImportRowStatus.IMPORTED.value
    row.error_message = None
    row.updated_at = datetime.now(UTC).replace(tzinfo=None)
    db.flush()
    return True


def _parse_excel_source_records(file_bytes: bytes) -> list[tuple[int, ConsumptionSourceRecord, dict[str, str | None]]]:
    try:
        workbook = load_workbook(filename=BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise PrepaidBillingValidationError("Excel 文件无法解析，请重新导出 .xlsx 模板后重试") from exc
    worksheet = workbook.worksheets[0]
    rows_iter = worksheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows_iter)]
    except StopIteration as exc:
        raise PrepaidBillingValidationError("Excel 文件缺少表头") from exc
    if headers != EXCEL_HEADERS:
        raise PrepaidBillingValidationError("Excel 表头不符合模板要求，请先下载最新模板")

    parsed_rows: list[tuple[int, ConsumptionSourceRecord, dict[str, str | None]]] = []
    for row_index, values in enumerate(rows_iter, start=2):
        raw_values = {
            EXCEL_HEADERS[index]: _cell_to_text(values[index]) if index < len(values) else None
            for index in range(len(EXCEL_HEADERS))
        }
        if all((value or "").strip() == "" for value in raw_values.values()):
            continue
        parsed_rows.append(
            (
                row_index,
                ConsumptionSourceRecord(
                    record_source="EXCEL",
                    supplier_record_no=raw_values["供应商流水号"],
                    consumed_at_text=raw_values["消费时间"],
                    user_phone=raw_values["用户手机号"],
                    service_code=raw_values["服务编码"],
                    supplier_name=raw_values["供应商名称"],
                    item_name=raw_values["消费项目"],
                    quantity_text=raw_values["消费数量"],
                    unit=raw_values["用量单位"],
                    amount_text=raw_values["实际消费金额"],
                    order_no=raw_values["关联订单号"],
                    remark=raw_values["备注"],
                ),
                raw_values,
            )
        )
        if len(parsed_rows) > MAX_IMPORT_ROWS:
            raise PrepaidBillingValidationError("单次最多导入 10000 条数据行")
    workbook.close()
    return parsed_rows


def _batch_summary_response(batch: ConsumptionImportBatch) -> ConsumptionImportBatchSummaryResponse:
    return ConsumptionImportBatchSummaryResponse(
        id=batch.id,
        batch_no=batch.batch_no,
        record_source=batch.record_source,
        file_name=batch.file_name,
        file_sha256=batch.file_sha256,
        status=batch.status,
        total_rows=batch.total_rows,
        valid_rows=batch.valid_rows,
        invalid_rows=batch.invalid_rows,
        imported_rows=batch.imported_rows,
        uploaded_by=batch.uploaded_by,
        confirmed_by=batch.confirmed_by,
        uploaded_at=_utc_value(batch.uploaded_at) or datetime.now(UTC),
        confirmed_at=_utc_value(batch.confirmed_at),
        expires_at=_utc_value(batch.expires_at) or datetime.now(UTC),
        error_message=batch.error_message,
    )


def _row_response(row: ConsumptionImportRow) -> ConsumptionImportRowResponse:
    return ConsumptionImportRowResponse(
        id=row.id,
        row_no=row.row_no,
        record_source=row.record_source,
        supplier_record_no=row.supplier_record_no,
        consumed_at=_utc_value(row.consumed_at),
        user_phone=row.user_phone,
        service_code=row.service_code,
        supplier_name=row.supplier_name,
        item_name=row.item_name,
        quantity=row.quantity,
        unit=row.unit,
        amount=row.amount,
        order_no=row.order_no,
        remark=row.remark,
        site_user_id=row.site_user_id,
        service_name=row.service_name,
        order_id=row.order_id,
        row_status=row.row_status,
        error_message=row.error_message,
        raw_values=_decode_raw_values(row.raw_values_json),
    )


def _consumption_record_response(
    record: ConsumptionRecord,
    user_name: str | None,
    user_phone_masked: str | None,
    order_no: str | None,
) -> ConsumptionRecordResponse:
    return ConsumptionRecordResponse(
        id=record.id,
        supplier_name=record.supplier_name,
        supplier_record_no=record.supplier_record_no,
        site_user_id=record.site_user_id,
        user_name=user_name,
        user_phone_masked=user_phone_masked,
        service_code=record.service_code,
        service_name=record.service_name,
        order_id=record.order_id,
        order_no=order_no,
        consumed_at=_utc_value(record.consumed_at) or datetime.now(UTC),
        item_name=record.item_name,
        quantity=record.quantity,
        unit=record.unit,
        amount=Decimal(record.amount).quantize(Decimal("0.01")),
        currency=record.currency,
        remark=record.remark,
        created_at=_utc_value(record.created_at) or datetime.now(UTC),
    )


def _expire_batch_if_needed(
    db: Session,
    batch: ConsumptionImportBatch,
    raise_conflict: bool,
) -> None:
    if batch.status != ConsumptionImportBatchStatus.PREVIEWED.value:
        return
    now = datetime.now(UTC).replace(tzinfo=None)
    if batch.expires_at > now:
        return
    batch.status = ConsumptionImportBatchStatus.EXPIRED.value
    batch.updated_at = now
    db.commit()
    if raise_conflict:
        raise PrepaidBillingConflictError("导入批次已过期，请重新上传文件")


def _generate_import_batch_no(now: datetime) -> str:
    utc_now = now.astimezone(UTC)
    suffix = "".join(secrets.choice(IMPORT_BATCH_ALPHABET) for _ in range(8))
    return f"CIB{utc_now.strftime('%Y%m%d%H%M%S')}{suffix}"


def _decode_raw_values(raw_values_json: str) -> dict[str, str | None]:
    data = json.loads(raw_values_json)
    return {str(key): (None if value is None else str(value)) for key, value in data.items()}


def _cell_to_text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value).strip()


def _normalize_text(value: str | None) -> str | None:
    normalized = (value or "").strip()
    if not normalized:
        return None
    return normalized


def _normalize_optional_text(value: str | None, max_length: int, field_label: str, errors: list[str]) -> str | None:
    normalized = _normalize_text(value)
    if normalized and len(normalized) > max_length:
        errors.append(f"{field_label}长度不能超过 {max_length}")
        return normalized[:max_length]
    return normalized or None


def _normalize_required_text(value: str | None, field_label: str, max_length: int, errors: list[str]) -> str | None:
    normalized = _normalize_text(value)
    if not normalized:
        errors.append(f"{field_label}不能为空")
        return None
    if len(normalized) > max_length:
        errors.append(f"{field_label}长度不能超过 {max_length}")
        return normalized[:max_length]
    return normalized


def _normalize_phone(value: str | None, errors: list[str]) -> str | None:
    normalized = _normalize_text(value)
    if not normalized:
        errors.append("用户手机号不能为空")
        return None
    if not PHONE_PATTERN.fullmatch(normalized):
        errors.append("用户手机号格式不合法")
    return normalized


def _normalize_code(value: str | None, field_label: str, errors: list[str]) -> str | None:
    normalized = _normalize_text(value)
    if not normalized:
        errors.append(f"{field_label}不能为空")
        return None
    normalized = normalized.upper()
    if not CODE_PATTERN.fullmatch(normalized):
        errors.append(f"{field_label}格式不合法")
    return normalized


def _normalize_amount(value: str | None, errors: list[str]) -> Decimal | None:
    normalized = _normalize_text(value)
    if not normalized:
        errors.append("实际消费金额不能为空")
        return None
    try:
        amount = Decimal(normalized)
    except InvalidOperation:
        errors.append("实际消费金额格式不合法")
        return None
    if amount <= 0:
        errors.append("实际消费金额必须大于 0")
        return None
    if amount.as_tuple().exponent < -2:
        errors.append("实际消费金额最多保留两位小数")
        return None
    return amount.quantize(Decimal("0.01"))


def _normalize_quantity(value: str | None, errors: list[str]) -> Decimal | None:
    normalized = _normalize_text(value)
    if not normalized:
        return None
    try:
        quantity = Decimal(normalized)
    except InvalidOperation:
        errors.append("消费数量格式不合法")
        return None
    if quantity < 0:
        errors.append("消费数量不能小于 0")
        return None
    if quantity.as_tuple().exponent < -6:
        errors.append("消费数量最多保留六位小数")
        return None
    return quantity.quantize(Decimal("0.000001"))


def _normalize_consumed_at(value: str | None, errors: list[str]) -> datetime | None:
    normalized = _normalize_text(value)
    if not normalized:
        errors.append("消费时间不能为空")
        return None
    try:
        local_value = datetime.strptime(normalized, "%Y-%m-%d %H:%M:%S").replace(tzinfo=SHANGHAI_TZ)
    except ValueError:
        errors.append("消费时间格式必须为 yyyy-MM-dd HH:mm:ss")
        return None
    return local_value.astimezone(UTC).replace(tzinfo=None)


def _mask_phone(phone: str) -> str:
    if len(phone) >= 7:
        return f"{phone[:3]}****{phone[-4:]}"
    return "****"


def _utc_value(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    return value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)


def _naive_utc(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is None:
        return value
    return value.astimezone(UTC).replace(tzinfo=None)
